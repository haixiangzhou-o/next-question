#!/usr/bin/env python3
"""从 Word / PPT 提取纯文本，作为 AI 结构化的输入。

支持格式:
  .docx   直接提取（段落 + 表格，按文档顺序）
  .pptx   逐页提取形状文本
  .doc    先调用 LibreOffice(soffice) 转为 .docx 再提取

用法:
  python extract.py <input> [-o output.txt]

示例:
  python extract.py "../raw_materials/2410自考《15043中国近现代史纲要》真题和答案.docx" -o 2410.txt
"""
import argparse
import shutil
import subprocess
import sys
from pathlib import Path


def iter_block_items(parent):
    """按文档顺序产出段落与表格。"""
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in parent.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def extract_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    lines = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                lines.append(text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for idx, slide in enumerate(prs.slides, 1):
        parts.append(f"===== 第 {idx} 页 =====")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    parts.append(" | ".join(cells))
    return "\n".join(parts)


def convert_doc_to_docx(path: Path) -> Path:
    """调用 LibreOffice 将 .doc 转为 .docx，返回转换后路径。"""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(
            "未找到 LibreOffice(soffice)。请安装 LibreOffice 后再处理 .doc 文件，"
            "或手动将 .doc 另存为 .docx。"
        )
    out_dir = path.parent
    subprocess.run(
        [soffice, "--headless", "--convert-to", "docx",
         "--outdir", str(out_dir), str(path)],
        check=True,
    )
    return out_dir / (path.stem + ".docx")


def main(argv=None):
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("input", help="输入文件（.doc / .docx / .pptx）")
    parser.add_argument("-o", "--output", help="输出文本文件路径（默认与输入同目录同名 .txt）")
    args = parser.parse_args(argv)

    src = Path(args.input)
    if not src.exists():
        print(f"文件不存在: {src}", file=sys.stderr)
        return 2

    suffix = src.suffix.lower()
    try:
        if suffix == ".doc":
            print(f"转换 .doc -> .docx: {src.name}")
            src = convert_doc_to_docx(src)
            suffix = ".docx"
        if suffix == ".docx":
            text = extract_docx(src)
        elif suffix == ".pptx":
            text = extract_pptx(src)
        else:
            print(f"不支持的文件格式: {suffix}", file=sys.stderr)
            return 2
    except Exception as e:
        print(f"提取失败: {e}", file=sys.stderr)
        return 1

    out = Path(args.output) if args.output else src.with_suffix(".txt")
    out.write_text(text, encoding="utf-8")
    print(f"已提取 {len(text)} 字符 -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
